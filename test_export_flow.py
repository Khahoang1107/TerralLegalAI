"""
Test end-to-end luồng xuất file DOCX từ biểu mẫu.
Chạy: python test_export_flow.py

Yêu cầu:
- Docker đang chạy (http://localhost:8000)
- Có file DOCX mẫu để upload (đặt tên TEST_DOCX bên dưới)
- Đã có tài khoản hoặc script sẽ tự đăng ký tài khoản mới
"""

import requests
import json
import os
import sys

BASE_URL = "http://localhost:8000/api/v1"

# ─── CẤU HÌNH ─────────────────────────────────────────────────────────────────
TEST_EMAIL    = "test_export@example.com"
TEST_PASSWORD = "TestPass123!"
TEST_NAME     = "Test Export User"

# Đường dẫn file DOCX mẫu để test — thay bằng file thật của bạn
# Nếu không có file, script sẽ tạo một file DOCX đơn giản tự động
TEST_DOCX = r"d:\TerraLegalAI\backend\data\templates\sample.docx"

# Dữ liệu điền vào form khi test export
TEST_PAYLOAD = {
    "ho_ten":    "Nguyễn Văn A",
    "cmnd":      "123456789012",
    "dia_chi":   "123 Đường Lê Lợi, Quận 1, TP.HCM",
    "dien_thoai": "0912345678",
    "noi_dung":  "Chuyển nhượng quyền sử dụng đất",
    "ngay":      "28",
    "thang":     "07",
    "nam":       "2026",
}
# ──────────────────────────────────────────────────────────────────────────────


def step(msg):
    print(f"\n{'='*60}")
    print(f"  {msg}")
    print('='*60)

def ok(msg):  print(f"  ✅ {msg}")
def err(msg): print(f"  ❌ {msg}"); sys.exit(1)
def info(msg): print(f"  ℹ️  {msg}")


def create_sample_docx(path: str):
    """Tạo file DOCX đơn giản có các vùng trống để test nếu chưa có file."""
    try:
        from docx import Document
        from docx.shared import Pt
        doc = Document()
        doc.add_paragraph("CỘNG HÒA XÃ HỘI CHỦ NGHĨA VIỆT NAM").runs[0].bold = True
        doc.add_paragraph("ĐƠN TEST XUẤT FILE")
        doc.add_paragraph("1. Họ và tên: {{ ho_ten }}")
        doc.add_paragraph("2. CMND/CCCD: {{ cmnd }}")
        doc.add_paragraph("3. Địa chỉ: {{ dia_chi }}")
        doc.add_paragraph("4. Điện thoại: {{ dien_thoai }}")
        doc.add_paragraph("5. Nội dung: {{ noi_dung }}")
        doc.add_paragraph("Ngày {{ ngay }} tháng {{ thang }} năm {{ nam }}")
        os.makedirs(os.path.dirname(path), exist_ok=True)
        doc.save(path)
        return True
    except Exception as e:
        print(f"  ⚠️  Không tạo được file DOCX mẫu: {e}")
        return False


def main():
    session = requests.Session()

    # ── BƯỚC 0: Chuẩn bị file DOCX ───────────────────────────────────────────
    step("BƯỚC 0: Chuẩn bị file DOCX mẫu")
    if not os.path.exists(TEST_DOCX):
        info(f"Không tìm thấy {TEST_DOCX}, đang tạo file mẫu đơn giản...")
        if create_sample_docx(TEST_DOCX):
            ok(f"Đã tạo file mẫu tại: {TEST_DOCX}")
        else:
            err("Không thể tạo file DOCX mẫu. Hãy đặt file DOCX vào TEST_DOCX và chạy lại.")
    else:
        ok(f"Dùng file: {TEST_DOCX}")

    # ── BƯỚC 1: Đăng ký tài khoản (nếu chưa có) ──────────────────────────────
    step("BƯỚC 1: Đăng ký / Đăng nhập")
    r = session.post(f"{BASE_URL}/auth/register", json={
        "full_name": TEST_NAME,
        "email":     TEST_EMAIL,
        "password":  TEST_PASSWORD,
    })
    if r.status_code in (200, 201):
        ok(f"Đăng ký thành công: {TEST_EMAIL}")
    elif r.status_code == 400 and "already" in r.text.lower():
        info("Tài khoản đã tồn tại, tiếp tục đăng nhập...")
    else:
        info(f"Đăng ký: {r.status_code} — {r.text[:100]}")

    # Đăng nhập
    r = session.post(f"{BASE_URL}/auth/login", json={
        "email":    TEST_EMAIL,
        "password": TEST_PASSWORD,
    })
    if r.status_code != 200:
        err(f"Đăng nhập thất bại: {r.status_code} — {r.text}")

    token = r.json().get("access_token")
    if not token:
        err("Không lấy được access_token!")
    ok(f"Đăng nhập thành công. Token: {token[:30]}...")
    session.headers.update({"Authorization": f"Bearer {token}"})

    # ── BƯỚC 2: Upload biểu mẫu ──────────────────────────────────────────────
    step("BƯỚC 2: Upload biểu mẫu DOCX + JSON schema")
    schema_json = json.dumps({
        "name": "Đơn Test Xuất File",
        "procedure_type": "test_export",
        "description": "Form test tự động",
        "fields": [
            {"key": "ho_ten",     "name": "Họ và tên",   "type": "string",  "required": True,  "description": "Họ và tên đầy đủ"},
            {"key": "cmnd",       "name": "CMND/CCCD",   "type": "string",  "required": True,  "description": "Số CMND hoặc CCCD"},
            {"key": "dia_chi",    "name": "Địa chỉ",     "type": "string",  "required": True,  "description": "Địa chỉ thường trú"},
            {"key": "dien_thoai", "name": "Điện thoại",  "type": "string",  "required": False, "description": "Số điện thoại liên hệ"},
            {"key": "noi_dung",   "name": "Nội dung",    "type": "string",  "required": True,  "description": "Nội dung biến động"},
            {"key": "ngay",       "name": "Ngày",        "type": "string",  "required": True,  "description": "Ngày ký đơn"},
            {"key": "thang",      "name": "Tháng",       "type": "string",  "required": True,  "description": "Tháng ký đơn"},
            {"key": "nam",        "name": "Năm",         "type": "string",  "required": True,  "description": "Năm ký đơn"},
        ]
    }, ensure_ascii=False)

    with open(TEST_DOCX, "rb") as f:
        r = session.post(f"{BASE_URL}/forms/upload", files={
            "json_file": ("schema.json", schema_json.encode("utf-8"), "application/json"),
            "docx_file": (os.path.basename(TEST_DOCX), f, "application/vnd.openxmlformats-officedocument.wordprocessingml.document"),
        })

    if r.status_code != 200:
        err(f"Upload thất bại: {r.status_code} — {r.text}")

    form_id = r.json().get("form_id")
    ok(f"Upload thành công! form_id = {form_id}")

    # Cần mô phỏng save-mapping để gán tag vào file docx.
    step("BƯỚC 2.5: Save mapping")
    r = session.post(f"{BASE_URL}/forms/{form_id}/save-mapping", json={
        "mapping": {
            "1": "ho_ten",
            "2": "cmnd",
            "3": "dia_chi",
            "4": "dien_thoai",
            "5": "noi_dung",
            "6": "ngay",
            "7": "thang",
            "8": "nam"
        }
    })
    
    if r.status_code == 200:
         ok("Save mapping thành công")
    else:
         print(f"Save mapping thất bại: {r.text}, nhưng nếu DOCX mẫu đã có sẵn tag {{{{ }}}} thì vẫn có thể tiếp tục")
         

    # ── BƯỚC 3: Export DOCX ──────────────────────────────────────────────────
    step("BƯỚC 3: Export DOCX với dữ liệu điền vào")
    info(f"Payload: {json.dumps(TEST_PAYLOAD, ensure_ascii=False)}")

    r = session.post(
        f"{BASE_URL}/forms/export/{form_id}",
        json=TEST_PAYLOAD,
        params={"format": "docx"},
    )

    if r.status_code != 200:
        err(f"Export thất bại: {r.status_code} — {r.text[:300]}")

    output_path = f"test_output_{form_id[:8]}.docx"
    with open(output_path, "wb") as f:
        f.write(r.content)
    ok(f"Export DOCX thành công!")
    ok(f"File lưu tại: {os.path.abspath(output_path)} ({len(r.content):,} bytes)")

    # ── BƯỚC 4: Xác minh nội dung file output ─────────────────────────────────
    step("BƯỚC 4: Xác minh nội dung file DOCX output")
    try:
        from docx import Document
        doc_out = Document(output_path)
        full_text = "\n".join(p.text for p in doc_out.paragraphs)

        checks = {
            "Nguyễn Văn A":    "ho_ten",
            "123456789012":    "cmnd",
            "123 Đường Lê Lợi": "dia_chi",
            "Chuyển nhượng":   "noi_dung",
            "28":              "ngay",
            "07":              "thang",
            "2026":            "nam",
        }
        all_ok = True
        for value, field in checks.items():
            if value in full_text:
                ok(f"'{value}' (field: {field}) — CÓ trong file output ✓")
            else:
                print(f"  ⚠️  '{value}' (field: {field}) — KHÔNG tìm thấy trong file output")
                all_ok = False

        # Kiểm tra không còn {{ tags }} thừa
        if "{{" in full_text:
            print(f"  ⚠️  Còn sót {{ tags }} chưa được render!")
            all_ok = False
        else:
            ok("Không còn {{ tags }} thừa trong file output ✓")

        if all_ok:
            print(f"\n{'='*60}")
            print("  🎉 TOÀN BỘ TEST PASSED — Luồng export hoạt động chính xác!")
            print(f"  📄 Mở file '{output_path}' trong Word để kiểm tra trực quan.")
            print('='*60)
        else:
            print(f"\n  ⚠️  Một số field chưa được render đúng. Kiểm tra lại template DOCX.")

    except Exception as e:
        print(f"  ⚠️  Không thể đọc file output để verify: {e}")
        ok(f"File đã được tạo. Hãy mở '{output_path}' bằng Word để kiểm tra thủ công.")

    # ── BƯỚC 5: Kiểm tra lịch sử submissions ─────────────────────────────────
    step("BƯỚC 5: Kiểm tra lịch sử submissions đã lưu vào DB")
    r = session.get(f"{BASE_URL}/forms/{form_id}/submissions")
    if r.status_code == 200:
        subs = r.json()
        ok(f"Tìm thấy {len(subs)} bản ghi submission trong DB")
        if subs:
            info(f"Submission mới nhất: format={subs[0].get('output_format')}, file={subs[0].get('output_filename')}")
    else:
        print(f"  ⚠️  Lấy submissions: {r.status_code} — {r.text[:100]}")


if __name__ == "__main__":
    main()
