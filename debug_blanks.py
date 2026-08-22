import docx, re
blank_pattern = re.compile(
    r'(?:[\._ ](?:&nbsp;|\s)*){3,}'
    r'|\t+'
    r'|[\u2610\u25a1]'
    r'|(?:\u2026){1,}'
    r'|(?:\u2025){1,}'
    r'|(?:[\u2013\u2014]){2,}'
)
d = docx.Document('backend/data/templates/83e2efc7-6e5d-46b0-8a62-e70aeb311d5e.docx')
idx = 1
out = []
def p(runs):
    global idx
    for r in runs:
        matches = list(blank_pattern.finditer(r.text))
        for m in matches:
            out.append(f'{idx}: text match in "{r.text}"')
            idx += 1

for pgh in d.paragraphs:
    p(pgh.runs)

for t in d.tables:
    for row in t.rows:
        for cell in row.cells:
            if not cell.text.strip():
                out.append(f'{idx}: empty cell')
                idx += 1
            else:
                for pgh in cell.paragraphs:
                    p(pgh.runs)

open('debug_blanks.txt', 'w', encoding='utf-8').write('\n'.join(out))
