import docx
from docx.oxml.shared import OxmlElement
from docx.oxml.ns import qn

doc = docx.Document()
p = doc.add_paragraph('Test ')
run = p.add_run('1')
rPr = run._r.get_or_add_rPr()
bdr = OxmlElement('w:bdr')
bdr.set(qn('w:val'), 'single')
bdr.set(qn('w:sz'), '4')
bdr.set(qn('w:space'), '0')
bdr.set(qn('w:color'), 'auto')
rPr.append(bdr)

doc.save('test_border.docx')
