import mammoth

docx_path = "../data/uploaded/test.docx" # mock path
style_map = """
p[alignment='center'] => p.text-center:fresh
p[alignment='right'] => p.text-right:fresh
"""

# Let's create a minimal docx in memory with a centered paragraph
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
import io

doc = Document()
p = doc.add_paragraph("CENTERED TEXT")
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p = doc.add_paragraph("RIGHT TEXT")
p.alignment = WD_ALIGN_PARAGRAPH.RIGHT

doc_io = io.BytesIO()
doc.save(doc_io)
doc_io.seek(0)

result = mammoth.convert_to_html(doc_io, style_map=style_map)
print("HTML:", result.value)
