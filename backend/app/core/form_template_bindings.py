"""Bind Word blanks and restore unbound signing dates in existing templates."""
import io
import re
import unicodedata

import docx


BLANK_PATTERN = re.compile(
    r'(?:[\._ ](?:&nbsp;|\s)*){3,}|\t+|[\u2610\u25a1]'
    r'|\u2026+|\u2025+|[\u2013\u2014]{2,}'
)
W_NAMESPACE = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"


def normalize_docx_symbols(document) -> int:
    """Convert proprietary Wingdings/Symbol XML elements (<w:sym>) into standard Unicode checkboxes.

    Without this, LibreOffice on Linux falls back to a generic missing glyph which squishes
    the checkbox horizontally (bị bóp chiều ngang).
    """
    from docx.oxml import OxmlElement
    converted = 0
    all_paragraphs = list(document.paragraphs)
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                all_paragraphs.extend(cell.paragraphs)

    for p in all_paragraphs:
        for r in p.runs:
            syms = r._r.findall(f"{{{W_NAMESPACE}}}sym")
            for sym in syms:
                font = sym.attrib.get(f"{{{W_NAMESPACE}}}font", "")
                char_hex = sym.attrib.get(f"{{{W_NAMESPACE}}}char", "").upper()
                if "wingdings" in font.lower() or "symbol" in font.lower():
                    rep = "☐"
                    if char_hex in ("F0A8", "F0FE", "F052", "F053", "F0A3"):
                        rep = "☑"
                    t_elem = OxmlElement("w:t")
                    t_elem.text = rep
                    r.font.name = "DejaVu Sans"
                    r._r.replace(sym, t_elem)
                    converted += 1
    return converted


def current_date_rule(field):
    """Return an explicit rule, or infer one for older current-date fields."""
    rule = field.get('auto_rule')
    if rule in {'current_date_day', 'current_date_month', 'current_date_year'}:
        return rule
    if field.get('value_source') != 'current_date':
        return None
    name = unicodedata.normalize('NFD', str(field.get('name') or ''))
    name = ''.join(char for char in name if unicodedata.category(char) != 'Mn').lower()
    if re.search(r'\b(thang|month)\b', name):
        return 'current_date_month'
    if re.search(r'\b(nam|year)\b', name):
        return 'current_date_year'
    if re.search(r'\b(ngay|day)\b', name):
        return 'current_date_day'
    return None


def _paragraphs(document, include_empty=False):
    yield from document.paragraphs
    seen = set()
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                if cell._tc in seen:
                    continue
                seen.add(cell._tc)
                if include_empty and not cell.text.strip():
                    paragraph = cell.paragraphs[0] if cell.paragraphs else cell.add_paragraph()
                    paragraph.add_run('   ')
                yield from cell.paragraphs


def _replace_spans(paragraph, changes):
    """Replace text spans across runs while retaining the surrounding formatting."""
    for start, end, replacement in sorted(changes, reverse=True):
        offset = 0
        inserted = False
        for run in paragraph.runs:
            text = run.text
            run_end = offset + len(text)
            if offset < end and run_end > start:
                left = max(0, start - offset)
                right = min(len(text), end - offset)
                run.text = text[:left] + (replacement if not inserted else '') + text[right:]
                inserted = True
            offset = run_end


def _restore_signing_date(document, fields):
    by_rule = {
        current_date_rule(field): field for field in fields
        if isinstance(field, dict) and field.get('key')
    }
    rules = {
        'ngày': 'current_date_day',
        'tháng': 'current_date_month',
        'năm': 'current_date_year',
    }
    date_fields = [by_rule.get(rule) for rule in rules.values()]
    if not any(date_fields):
        return

    locality = None
    day = by_rule.get('current_date_day') or {}
    day_zones = day.get('visual_zones') or []
    if day_zones:
        dz = day_zones[0]
        candidates = []
        for field in fields:
            if not isinstance(field, dict) or field in date_fields:
                continue
            for zone in field.get('visual_zones') or []:
                if zone.get('page') != dz.get('page'):
                    continue
                if zone.get('x', 0) >= dz.get('x', 0):
                    continue
                if abs(zone.get('y', 0) - dz.get('y', 0)) <= max(zone.get('h', 0), dz.get('h', 0), 5):
                    candidates.append((zone.get('x', 0), field))
        if candidates:
            locality = max(candidates, key=lambda item: item[0])[1]
    if locality is None:
        locality = next((
            field for field in fields if isinstance(field, dict)
            and 'nơi' in field.get('name', '').lower()
            and 'đơn' in field.get('name', '').lower()
            and any(word in field.get('name', '').lower() for word in ('tỉnh', 'thành', 'địa danh'))
        ), None)

    for paragraph in _paragraphs(document):
        text = paragraph.text
        if not all(word in text for word in rules):
            continue
        # Signing lines start with a locality blank/binding or directly with
        # "ngày"; dates embedded in legal prose or personal details are excluded.
        if not re.match(r'^\s*(?:(?:[.\u2026\u2025_]+|\{\{\s*\w+\s*\}\})\s*,?\s*)?ngày\b', text):
            continue
        changes = []
        for word, rule in rules.items():
            field = by_rule.get(rule)
            if not field:
                continue
            # Existing Jinja bindings remain intact; only restore literal blanks.
            match = re.search(r'\b' + word + r'\s*([.\u2026\u2025_]+)', text)
            if match:
                start, end = match.span(1)
                changes.append((start, end, '{{ ' + field['key'] + ' }}'))
        if locality and changes:
            prefix = re.search(r'([.\u2026\u2025_]+)\s*,\s*ngày\b', text)
            if prefix:
                changes.append((*prefix.span(1), '{{ ' + locality['key'] + ' }}'))
        _replace_spans(paragraph, changes)


def bound_template_stream(template_path, mapping, fields=None):
    """Render on a copy, preserving the uploaded template and its existing tags."""
    document = docx.Document(template_path)
    normalize_docx_symbols(document)
    if fields is not None:
        _restore_signing_date(document, fields)
    if mapping:
        # Check if this document already has tags from mapping.
        # If the template was already injected with Jinja tags, re-applying
        # mapping by blank-index will recount only the remaining unmapped blanks,
        # shifting indices and corrupting unmapped lines (e.g. turning the 2nd line
        # of 'Nội dung biến động' into {{ field_9003 }} 'Nơi tiếp nhận hồ sơ').
        all_text = " ".join(p.text for p in _paragraphs(document, include_empty=False))
        existing_tags = set(re.findall(r'\{\{\s*(\w+)\s*\}\}', all_text))
        mapped_keys = set(mapping.values()) if mapping else set()
        already_bound = bool(mapped_keys and len(mapped_keys & existing_tags) >= 1)

        if not already_bound:
            # Initial binding uses the same single table pass as blank detection.
            index = 0
            for paragraph in _paragraphs(document, include_empty=True):
                for run in paragraph.runs:
                    matches = list(BLANK_PATTERN.finditer(run.text))
                    changes = []
                    for match in matches:
                        index += 1
                        key = mapping.get(str(index))
                        if key and re.fullmatch(r'\w+', key):
                            replacement = (
                                f'{{% if {key} %}}☑{{% else %}}☐{{% endif %}}'
                                if match.group() in ('☐', '□') else f'{{{{ {key} }}}}'
                            )
                            start, end = match.span()
                            matched = match.group()
                            if matched.strip():
                                start += len(matched) - len(matched.lstrip())
                                end -= len(matched) - len(matched.rstrip())
                            changes.append((start, end, replacement))
                    text = run.text
                    for start, end, replacement in reversed(changes):
                        text = text[:start] + replacement + text[end:]
                    if changes:
                        run.text = text

    # Self-healing: if a header tag like "Kính gửi" accidentally appears
    # inside "2. Nội dung biến động" (from earlier blank-shifting on disk), remove it.
    kinh_gui_tag = None
    for p in document.paragraphs:
        if 'kính gửi' in p.text.lower():
            m = re.search(r'\{\{\s*(\w+)\s*\}\}', p.text)
            if m:
                kinh_gui_tag = m.group(1)
                break
    if kinh_gui_tag:
        in_noi_dung = False
        for p in document.paragraphs:
            p_lower = p.text.lower()
            if 'nội dung biến động' in p_lower:
                in_noi_dung = True
                continue
            if in_noi_dung:
                if any(h in p_lower for h in ['giấy tờ liên quan', '3.', 'cam đoan']):
                    in_noi_dung = False
                    continue
                if kinh_gui_tag in p.text:
                    for r in p.runs:
                        if kinh_gui_tag in r.text:
                            r.text = re.sub(r'\{\{\s*' + kinh_gui_tag + r'\s*\}\}', '', r.text)

    stream = io.BytesIO()
    document.save(stream)
    stream.seek(0)
    return stream
